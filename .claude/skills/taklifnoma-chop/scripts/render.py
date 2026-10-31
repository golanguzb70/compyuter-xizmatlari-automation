#!/usr/bin/env python3
"""Render an invitation text to a print-ready Word (.docx) file at exact card size.

The shop hands out ONLY .docx files. --preview additionally draws a PNG into
a temp folder (never next to the order) so the layout can be checked by eye.

Text file markup (one printed line per text line):
    # АССАЛОМУ АЛАЙКУМ!      title line: larger, bold
    ## ЖУМА   20   НОЯБР     big line: largest, bold (date blocks)
    {ЖАСУРЖОН}               name: red + bold
    **НАҲОРГИ ОШ**           bold
    (empty line)             half-line gap

The font size is fitted automatically: the largest size at which every line
fits the text area width and the whole block fits its height. Text is
centred horizontally and vertically inside the area.

Examples:
    render.py matn.txt --out buyurtmalar/2026-11-15-karimov/taklifnoma-1
    render.py matn.txt --preset a6-yotiq --guides --out test   # test print with outlines
    render.py matn.txt --card 21x10 --area 11x6 --offset 8.5,2 --out k
    render.py matn.txt --preset matn-11x6 --sheet A4 --out k   # several cards on A4, cut lines
    render.py matn.txt --preset maxsus-15x21 --frame gold --background fon.jpg --out maxsus
"""
import argparse
import copy
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Pt, RGBColor
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

HERE = os.path.dirname(os.path.abspath(__file__))
SKILL_DIR = os.path.dirname(HERE)

SCALE = {"body": 1.0, "title": 1.25, "big": 1.6}
EMPTY_LINE = 0.5  # fraction of a body line
SHEETS = {"A4": (21.0 * cm, 29.7 * cm), "A3": (29.7 * cm, 42.0 * cm)}
FRAME_COLORS = {"gold": "B8860B", "red": "9B1B30", "green": "1F5E3B"}


def emu(pt):
    return Emu(int(round(pt * 12700)))


def twips(pt):
    return str(int(round(pt * 20)))


def load_json(name):
    with open(os.path.join(SKILL_DIR, name), encoding="utf-8") as f:
        return json.load(f)


def register_font(key):
    fonts = load_json("fonts.json")
    if key not in fonts:
        sys.exit(f"Shrift «{key}» fonts.json da yo'q. Mavjud: {', '.join(fonts)}")
    spec = fonts[key]
    for style in ("regular", "bold"):
        if not os.path.exists(spec[style]):
            sys.exit(f"Shrift fayli topilmadi: {spec[style]}")
        pdfmetrics.registerFont(TTFont(f"{key}-{style}", spec[style]))
    return spec


def parse(text):
    """-> list of lines; each line = (kind, [(text, bold, is_name), ...]) or None for gaps."""
    lines = []
    for raw in text.rstrip("\n").splitlines():
        if not raw.strip():
            lines.append(None)
            continue
        kind = "body"
        if raw.startswith("## "):
            kind, raw = "big", raw[3:]
        elif raw.startswith("# "):
            kind, raw = "title", raw[2:]
        runs = []
        bold_default = kind != "body"
        for tok in re.split(r"(\{[^}]*\}|\*\*[^*]*\*\*)", raw):
            if not tok:
                continue
            if tok.startswith("{"):
                runs.append((tok[1:-1], True, True))
            elif tok.startswith("**"):
                runs.append((tok[2:-2], True, False))
            else:
                runs.append((tok, bold_default, False))
        lines.append((kind, runs))
    return lines


def line_width(runs, size, font_key):
    return sum(pdfmetrics.stringWidth(t, f"{font_key}-{'bold' if b else 'regular'}", size) for t, b, _ in runs)


def layout(lines, size, leading, font_key):
    widths, heights = [], []
    for ln in lines:
        if ln is None:
            heights.append(size * leading * EMPTY_LINE)
            widths.append(0)
        else:
            s = size * SCALE[ln[0]]
            heights.append(s * leading)
            widths.append(line_width(ln[1], s, font_key))
    return widths, heights


def fit_size(lines, area_w, area_h, leading, font_key, max_size, min_size):
    lo, hi = 1.0, max_size
    for _ in range(40):
        mid = (lo + hi) / 2
        w, h = layout(lines, mid, leading, font_key)
        if max(w) <= area_w and sum(h) <= area_h:
            lo = mid
        else:
            hi = mid
    size = round(lo - 0.05, 1)
    return size, size >= min_size


# ---------------------------------------------------------------- DOCX helpers

# Schema order of the children we touch; Word rejects files with elements out of order.
SUCCESSORS = {
    "w:pgBorders": ["w:lnNumType", "w:pgNumType", "w:cols", "w:formProt", "w:vAlign", "w:noEndnote",
                    "w:titlePg", "w:textDirection", "w:bidi", "w:rtlGutter", "w:docGrid", "w:printerSettings"],
    "w:tblBorders": ["w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription"],
    "w:tblLayout": ["w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription"],
    "w:tblW": ["w:jc", "w:tblCellSpacing", "w:tblInd", "w:tblBorders", "w:shd", "w:tblLayout", "w:tblCellMar",
               "w:tblLook"],
    "w:tcMar": ["w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark"],
    "w:tcW": ["w:gridSpan", "w:vMerge", "w:tcBorders", "w:shd", "w:noWrap", "w:tcMar", "w:textDirection",
              "w:tcFitText", "w:vAlign", "w:hideMark"],
}


def set_el(parent, tag, **attrs):
    el = parent.find(qn(tag))
    if el is None:
        el = OxmlElement(tag)
        after = [parent.find(qn(t)) for t in SUCCESSORS.get(tag, [])]
        after = [a for a in after if a is not None]
        if after:
            after[0].addprevious(el)
        else:
            parent.append(el)
    for k, v in attrs.items():
        el.set(qn(k), v)
    return el


def border_xml(parent, tag, sides, val, sz_eighths, color, space="0"):
    box = set_el(parent, tag)
    for side in sides:
        set_el(box, f"w:{side}", **{"w:val": val, "w:sz": str(sz_eighths), "w:space": space, "w:color": color})
    return box


def write_lines(container, lines, size, args, font_spec):
    """Append the invitation paragraphs to a document body or a table cell.

    The block is centred vertically with space before the first line (not
    vAlign), so it lands in the same place in Word, LibreOffice and viewers.
    """
    _, heights = layout(lines, size, args.leading, args.font)
    top_pad = max(0.0, (args.area_h - sum(heights)) / 2)
    first = first_line = True
    for ln in lines:
        if first and container.paragraphs:
            p = container.paragraphs[0]
        else:
            p = container.add_paragraph()
        first = False
        pf = p.paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.space_before = Pt(top_pad) if first_line else Pt(0)
        pf.space_after = Pt(0)
        first_line = False
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        if ln is None:
            pf.line_spacing = Pt(size * args.leading * EMPTY_LINE)
            continue
        s = size * SCALE[ln[0]]
        pf.line_spacing = Pt(s * args.leading)
        for t, b, is_name in ln[1]:
            r = p.add_run(t)
            r.font.name = font_spec["docx_name"]
            r._element.rPr.rFonts.set(qn("w:cs"), font_spec["docx_name"])
            r._element.rPr.rFonts.set(qn("w:hAnsi"), font_spec["docx_name"])
            r.font.size = Pt(s)
            r.font.bold = b
            r.font.italic = font_spec.get("italic", False)
            r.font.color.rgb = RGBColor.from_string((args.name_color if is_name else args.text_color).lstrip("#").upper())
    return container.paragraphs


def add_background(paragraph, image, w, h, x=0.0, y=0.0):
    """Insert a picture behind the text at (x, y) points from the page's top-left corner."""
    run = paragraph.add_run()
    inline = run.add_picture(image, width=emu(w), height=emu(h))._inline
    anchor = OxmlElement("wp:anchor")
    for k, v in {"distT": "0", "distB": "0", "distL": "0", "distR": "0", "simplePos": "0",
                 "relativeHeight": "0", "behindDoc": "1", "locked": "1", "layoutInCell": "1",
                 "allowOverlap": "1"}.items():
        anchor.set(k, v)
    sp = OxmlElement("wp:simplePos")
    sp.set("x", "0")
    sp.set("y", "0")
    anchor.append(sp)
    for axis, value in (("H", x), ("V", y)):
        pos = OxmlElement(f"wp:position{axis}")
        pos.set("relativeFrom", "page")
        off = OxmlElement("wp:posOffset")
        off.text = str(int(round(value * 12700)))
        pos.append(off)
        anchor.append(pos)
    anchor.append(copy.deepcopy(inline.find(qn("wp:extent"))))
    ee = OxmlElement("wp:effectExtent")
    for k in ("l", "t", "r", "b"):
        ee.set(k, "0")
    anchor.append(ee)
    anchor.append(OxmlElement("wp:wrapNone"))
    anchor.append(copy.deepcopy(inline.find(qn("wp:docPr"))))
    anchor.append(OxmlElement("wp:cNvGraphicFramePr"))
    anchor.append(copy.deepcopy(inline.find(qn("a:graphic"))))
    inline.getparent().replace(inline, anchor)


def new_document(page_w, page_h, left, right, top, bottom):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = emu(page_w), emu(page_h)
    sec.left_margin, sec.right_margin = emu(left), emu(right)
    sec.top_margin, sec.bottom_margin = emu(top), emu(bottom)
    sec.header_distance = sec.footer_distance = 0
    return doc, sec


def write_single(path, args, lines, size, font_spec):
    a = args
    left, top = a.card_x + a.off_x, a.card_y + a.off_y
    doc, sec = new_document(a.page_w, a.page_h, left, a.page_w - left - a.area_w,
                            top, a.page_h - top - a.area_h)
    if a.frame and a.page_is_card:
        pg = border_xml(sec._sectPr, "w:pgBorders", ["top", "left", "bottom", "right"],
                        "double", 12, a.frame_hex, space="14")
        pg.set(qn("w:offsetFrom"), "page")
    if a.guides:
        # dashed outline of the text area — test prints only
        pg = border_xml(sec._sectPr, "w:pgBorders", ["top", "left", "bottom", "right"],
                        "dashed", 4, "00A0E0", space="0")
        pg.set(qn("w:offsetFrom"), "text")
    paras = write_lines(doc, lines, size, a, font_spec)
    if a.background:
        add_background(paras[0], a.background, a.card_w, a.card_h, a.card_x, a.card_y)
    doc.save(path)


def write_sheet(path, args, lines, size, font_spec):
    a = args
    pw, ph = SHEETS[a.sheet]
    margin = 1.0 * cm
    best = None
    for W, H in [(pw, ph), (ph, pw)]:
        cols, rows = int((W - 2 * margin) // a.card_w), int((H - 2 * margin) // a.card_h)
        if cols * rows and (best is None or cols * rows > best[2] * best[3]):
            best = (W, H, cols, rows)
    if not best:
        sys.exit("Karta varaqqa sig'maydi")
    W, H, cols, rows = best
    gw, gh = cols * a.card_w, rows * a.card_h
    mx, my = (W - gw) / 2, (H - gh) / 2
    doc, sec = new_document(W, H, mx, mx, my, my)
    if W > H:
        from docx.enum.section import WD_ORIENT
        sec.orientation = WD_ORIENT.LANDSCAPE
    tbl = doc.add_table(rows=rows, cols=cols)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = tbl._tbl.tblPr
    set_el(tblPr, "w:tblLayout", **{"w:type": "fixed"})
    set_el(tblPr, "w:tblW", **{"w:w": twips(gw), "w:type": "dxa"})
    cut = ("double", 8, a.frame_hex) if a.frame else ("dotted", 4, "A0A0A0")
    border_xml(tblPr, "w:tblBorders", ["top", "left", "bottom", "right", "insideH", "insideV"], *cut)
    grid = tbl._tbl.tblGrid
    for gc in grid.findall(qn("w:gridCol")):
        gc.set(qn("w:w"), twips(a.card_w))
    pad_r = a.card_w - a.off_x - a.area_w
    pad_b = a.card_h - a.off_y - a.area_h
    for row in tbl.rows:
        trPr = row._tr.get_or_add_trPr()
        set_el(trPr, "w:trHeight", **{"w:val": twips(a.card_h), "w:hRule": "exact"})
        set_el(trPr, "w:cantSplit")
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            set_el(tcPr, "w:tcW", **{"w:w": twips(a.card_w), "w:type": "dxa"})
            mar = set_el(tcPr, "w:tcMar")
            for side, v in (("top", a.off_y), ("left", a.off_x), ("bottom", pad_b), ("right", pad_r)):
                set_el(mar, f"w:{side}", **{"w:w": twips(v), "w:type": "dxa"})
            write_lines(cell, lines, size, a, font_spec)
    # Word insists on a paragraph after a table; make it invisible
    last = doc.add_paragraph()
    last.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    last.paragraph_format.line_spacing = Pt(1)
    doc.save(path)
    print(f"Varaq {a.sheet}: {cols}×{rows} = {cols * rows} dona")


# ------------------------------------------------------- preview (temp only)

def write_preview(args, lines, size, colors_hex):
    """Draw the card as PNG in a temp folder so the layout can be checked by eye."""
    from reportlab.lib.colors import HexColor
    from reportlab.pdfgen import canvas

    tmp = tempfile.mkdtemp(prefix="taklifnoma-preview-")
    base = os.path.join(tmp, os.path.basename(args.out))
    c = canvas.Canvas(base + ".pdf", pagesize=(args.page_w, args.page_h))
    cx, cy = args.card_x, args.page_h - args.card_y - args.card_h  # card bottom-left, PDF coords
    if not args.page_is_card:
        c.setStrokeColor(HexColor("#999999"))
        c.setDash(4, 3)
        c.rect(cx, cy, args.card_w, args.card_h)
        c.setDash()
    if args.background:
        c.drawImage(args.background, cx, cy, args.card_w, args.card_h, preserveAspectRatio=False, mask="auto")
    if args.frame:
        c.setStrokeColor(HexColor("#" + args.frame_hex))
        m = 0.5 * cm
        c.setLineWidth(1.2)
        c.rect(m, m, args.card_w - 2 * m, args.card_h - 2 * m)
        c.setLineWidth(0.5)
        c.rect(m + 2, m + 2, args.card_w - 2 * m - 4, args.card_h - 2 * m - 4)
    ax, ay = cx + args.off_x, cy + args.card_h - args.off_y - args.area_h
    if args.guides:
        c.setStrokeColor(HexColor("#00A0E0"))
        c.setDash(3, 2)
        c.rect(ax, ay, args.area_w, args.area_h)
        c.setDash()
    widths, heights = layout(lines, size, args.leading, args.font)
    top = ay + args.area_h - (args.area_h - sum(heights)) / 2
    for ln, w, h in zip(lines, widths, heights):
        if ln is not None:
            s = size * SCALE[ln[0]]
            baseline = top - h + (h - s) / 2 + s * 0.22
            x = ax + (args.area_w - w) / 2
            for t, b, is_name in ln[1]:
                fname = f"{args.font}-{'bold' if b else 'regular'}"
                c.setFont(fname, s)
                c.setFillColor(HexColor(colors_hex["name"] if is_name else colors_hex["text"]))
                c.drawString(x, baseline, t)
                x += pdfmetrics.stringWidth(t, fname, s)
        top -= h
    c.showPage()
    c.save()
    png = base + ".png"
    if shutil.which("sips"):
        subprocess.run(["sips", "-s", "format", "png", base + ".pdf", "--out", png], check=True, capture_output=True)
    elif shutil.which("pdftoppm"):
        subprocess.run(["pdftoppm", "-png", "-r", "200", "-singlefile", base + ".pdf", base], check=True)
    else:
        return None
    os.remove(base + ".pdf")
    return png


def parse_pair(s, sep="x"):
    a, b = s.lower().replace("×", "x").split(sep)
    return float(a.replace(",", ".")) * cm, float(b.replace(",", ".")) * cm


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("text_file")
    ap.add_argument("--out", required=True, help="output path without extension (.docx is added)")
    ap.add_argument("--preset", help="card preset from cards.json")
    ap.add_argument("--card", help="card size WxH cm, e.g. 21x10")
    ap.add_argument("--area", help="text area WxH cm (default 11x6)")
    ap.add_argument("--offset", help="text area top-left from card top-left X,Y cm (default: centred)")
    ap.add_argument("--page", choices=list(SHEETS) + ["card"],
                    help="paper size set in the printer: A4 (card fed at --card-pos) or card (page = card)")
    ap.add_argument("--card-pos", help="card's top-left corner on the page X,Y cm (default: centred at the top edge)")
    ap.add_argument("--no-title", action="store_true",
                    help="drop '# ' title lines (card already has АССАЛОМУ АЛАЙКУМ printed)")
    ap.add_argument("--font", default="times-italic", help="key from fonts.json")
    ap.add_argument("--name-color", default="#C00000")
    ap.add_argument("--text-color", default="#000000")
    ap.add_argument("--frame", nargs="?", const="gold", help="decorative double frame: gold | red | green | #hex")
    ap.add_argument("--background", help="full-card background image (special designs)")
    ap.add_argument("--padding", type=float, default=0.2, help="inner safety margin inside the text area, cm")
    ap.add_argument("--leading", type=float, default=1.18)
    ap.add_argument("--max-size", type=float, default=16)
    ap.add_argument("--min-size", type=float, default=8)
    ap.add_argument("--size", type=float, help="force body font size (pt) instead of auto-fit")
    ap.add_argument("--guides", action="store_true", help="outline the text area (test print only!)")
    ap.add_argument("--sheet", choices=list(SHEETS), help="tile copies on a sheet with cut lines")
    ap.add_argument("--preview", action="store_true", help="also draw a PNG preview into a temp folder")
    args = ap.parse_args()

    preset = {}
    if args.preset:
        presets = load_json("cards.json")
        if args.preset not in presets:
            sys.exit(f"Preset «{args.preset}» yo'q. Mavjud: {', '.join(k for k in presets if not k.startswith('_'))}")
        preset = presets[args.preset]
    area = args.area or (f"{preset['area'][0]}x{preset['area'][1]}" if preset else "11x6")
    args.area_w, args.area_h = parse_pair(area)
    card = args.card or (f"{preset['card'][0]}x{preset['card'][1]}" if preset else area)
    args.card_w, args.card_h = parse_pair(card)
    offset = args.offset or (f"{preset['offset'][0]},{preset['offset'][1]}" if preset.get("offset") else None)
    if offset:
        args.off_x, args.off_y = parse_pair(offset, ",")
    else:
        args.off_x, args.off_y = (args.card_w - args.area_w) / 2, (args.card_h - args.area_h) / 2
    if args.off_x < 0 or args.off_y < 0 or args.off_x + args.area_w > args.card_w + 0.01 \
            or args.off_y + args.area_h > args.card_h + 0.01:
        sys.exit("Matn maydoni karta chegarasidan chiqib ketyapti — o'lchamlarni tekshiring")
    if args.background and args.sheet:
        sys.exit("--background faqat bitta karta uchun ishlaydi (--sheet bilan emas)")

    page = args.page or preset.get("page", "A4")
    if args.sheet:
        page = "card"
    args.page_is_card = page == "card"
    if args.page_is_card:
        args.page_w, args.page_h, args.card_x, args.card_y = args.card_w, args.card_h, 0.0, 0.0
    else:
        args.page_w, args.page_h = SHEETS[page]
        pos = args.card_pos or (f"{preset['card_pos'][0]},{preset['card_pos'][1]}" if preset.get("card_pos") else None)
        if pos:
            args.card_x, args.card_y = parse_pair(pos, ",")
        else:
            args.card_x, args.card_y = (args.page_w - args.card_w) / 2, 0.0
            print("⚠️  card_pos berilmagan — karta varaq tepasida markazda deb olindi. kalibrovka.py bilan o'lchang!")
        if args.card_x + args.card_w > args.page_w + 0.01 or args.card_y + args.card_h > args.page_h + 0.01:
            sys.exit(f"Karta {page} varaqdan chiqib ketyapti — card_pos ni tekshiring")
        if args.frame:
            print("⚠️  --frame faqat sahifa = karta bo'lganda ishlaydi (A4 da o'tkazib yuborildi)")
    no_title = args.no_title or preset.get("preprinted_title", False)
    args.frame_hex = (FRAME_COLORS.get(args.frame) or (args.frame or "#B8860B")).lstrip("#").upper()

    font_spec = register_font(args.font)
    with open(args.text_file, encoding="utf-8") as f:
        lines = parse(f.read())
    if no_title:
        dropped = [ln for ln in lines if ln and ln[0] == "title"]
        lines = [ln for ln in lines if not (ln and ln[0] == "title")]
        if dropped:
            print("ℹ️  Sarlavha qatori chiqarilmadi (kartada oldindan bosilgan)")
    fit_w, fit_h = args.area_w - 2 * args.padding * cm, args.area_h - 2 * args.padding * cm
    if args.size:
        size, ok = args.size, True
        w, h = layout(lines, size, args.leading, args.font)
        if max(w) > fit_w or sum(h) > fit_h:
            print("⚠️  Berilgan o'lchamda matn maydonga sig'maydi!", file=sys.stderr)
    else:
        size, ok = fit_size(lines, fit_w, fit_h, args.leading, args.font, args.max_size, args.min_size)
    if not ok:
        sys.exit(f"❌ Matn sig'maydi: shrift {size}pt bo'lib qoladi (minimum {args.min_size}pt). "
                 "Matnni qisqartiring yoki qatorlarni qayta bo'ling.")

    os.makedirs(os.path.dirname(os.path.abspath(args.out)) or ".", exist_ok=True)
    path = args.out + ".docx"
    (write_sheet if args.sheet else write_single)(path, args, lines, size, font_spec)
    widths, _ = layout(lines, size, args.leading, args.font)
    widest = max(range(len(widths)), key=lambda i: widths[i])
    print(f"DOCX: {path}")
    if not args.page_is_card:
        print(f"Sahifa {page}; karta varaqda X={args.card_x / cm:.1f}, Y={args.card_y / cm:.1f} sm dan; "
              f"matn chap-yuqori burchagi X={(args.card_x + args.off_x) / cm:.1f}, "
              f"Y={(args.card_y + args.off_y) / cm:.1f} sm")
    print(f"Karta {args.card_w / cm:.1f}×{args.card_h / cm:.1f} sm, matn maydoni "
          f"{args.area_w / cm:.1f}×{args.area_h / cm:.1f} sm, asosiy shrift {size}pt "
          f"(eng uzun qator: {widest + 1}-qator)")
    if args.guides:
        print("⚠️  --guides yoqilgan: bu fayl faqat sinov uchun, kartaga chiqarmang")
    if args.preview:
        png = write_preview(args, lines, size, {"name": args.name_color, "text": args.text_color})
        print(f"Ko'rinish (faqat tekshirish uchun, mijozga berilmaydi): {png}" if png
              else "Ko'rinish yaratib bo'lmadi (sips/pdftoppm yo'q)")


if __name__ == "__main__":
    main()
